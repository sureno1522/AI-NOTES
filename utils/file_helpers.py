import json
import os
from pathlib import Path
import pdfplumber
import docx
from pptx import Presentation
from .config import SUPPORTED_EXTENSIONS


def load_pdf_text(file_path: str) -> str:
    text = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text.append(page_text)
    return "\n\n".join(text)


def load_docx_text(file_path: str) -> str:
    doc = docx.Document(file_path)
    return "\n\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())


def load_pptx_text(file_path: str) -> str:
    presentation = Presentation(file_path)
    slides = []
    for slide in presentation.slides:
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text)
        slides.append("\n".join(slide_text))
    return "\n\n".join(slides)


def load_txt_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        return file.read()


def load_document(file_path: str) -> dict:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        content = load_pdf_text(file_path)
    elif ext == ".docx":
        content = load_docx_text(file_path)
    elif ext == ".pptx":
        content = load_pptx_text(file_path)
    elif ext == ".txt":
        content = load_txt_text(file_path)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")
    return {"path": file_path, "content": content, "metadata": {"source": os.path.basename(file_path), "type": ext}}


def split_text_with_overlap(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    text = text.strip()
    if not text:
        return []

    chunks = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = min(start + chunk_size, text_length)
        chunk = text[start:end]
        chunks.append(chunk)
        if end == text_length:
            break
        start = max(0, end - chunk_overlap)

    return chunks


def build_text_chunks(documents: list[dict], chunk_size: int = 1000, chunk_overlap: int = 200) -> list[dict]:
    chunks = []
    for document in documents:
        segments = split_text_with_overlap(document["content"], chunk_size, chunk_overlap)
        for index, chunk in enumerate(segments):
            chunks.append({
                "text": chunk,
                "metadata": {
                    "source": document["metadata"]["source"],
                    "type": document["metadata"]["type"],
                    "chunk": index,
                },
            })
    return chunks
